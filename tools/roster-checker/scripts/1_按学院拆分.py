# -*- coding: utf-8 -*-
"""从 2026 级团员花名册（按学院分段的大表）里抽两个学院的名单，各存一个 xlsx"""
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import conf

SRC = conf.roster_file()
OUT_DIR = conf.OUT_DIR

# 表格里的分段名 -> 输出文件名
TARGETS = {
    "人工智能学院": "01-人工智能学院（数据学院）-2026级团员花名册.xlsx",
    "机械学院": "02-机械学院-2026级团员花名册.xlsx",
}

HEAD = ["序号", "学号", "专业班级", "姓名", "籍贯", "性别", "出生年月", "入团时间", "入团地点", "备注"]
WIDTH = [6, 14, 26, 10, 20, 6, 10, 10, 34, 12]

wb_in = openpyxl.load_workbook(SRC, data_only=True)
ws_in = wb_in.active

# ---------- 1. 按"只有第一列有值"的行切段 ----------
sections, cur, buf = {}, None, []
for row in ws_in.iter_rows(min_row=1, values_only=True):
    a = row[0]
    rest = [c for c in row[1:] if c not in (None, "")]
    is_title = a not in (None, "") and not rest
    if is_title:
        if cur:
            sections[cur] = buf
        name = str(a).strip()
        cur = None if name.startswith("合计") else name
        buf = []
        continue
    if cur is None:
        continue
    if str(a).strip() == "序号":      # 表头行
        continue
    if not row[3]:                    # 没姓名，跳过空行
        continue
    buf.append(list(row[:10]))
if cur:
    sections[cur] = buf

print("识别到的分段：")
for k, v in sections.items():
    print(f"  {k:8} {len(v):4} 人")

# ---------- 2. 导出 ----------
os.makedirs(OUT_DIR, exist_ok=True)
thin = Side(style="thin", color="D0D0D0")
for key, fname in TARGETS.items():
    rows = sections.get(key, [])
    if not rows:
        print(f"⚠️ 没找到「{key}」")
        continue

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = key[:31]

    # 标题
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    c = ws.cell(row=1, column=1, value=f"团员花名册（2026级）—— {key}　　共 {len(rows)} 人")
    c.font = Font(size=14, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # 表头
    for j, h in enumerate(HEAD, 1):
        cell = ws.cell(row=2, column=j, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[2].height = 22

    # 数据
    for i, r in enumerate(rows, 3):
        for j, v in enumerate(r, 1):
            if j == 2 and isinstance(v, (int, float)):     # 学号防丢精度
                v = str(int(v))
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(
                horizontal="center" if j in (1, 6, 7, 8) else "left", vertical="center")
            if j == 2:
                cell.number_format = "@"

    for j, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:J{len(rows) + 2}"

    path = os.path.join(OUT_DIR, fname)
    wb.save(path)
    print(f"写出 {fname}  {len(rows)} 人  {os.path.getsize(path) // 1024} KB")

print("\n输出目录:", OUT_DIR)
