# -*- coding: utf-8 -*-
"""把统计表和问题清单按学院拆成两份，分别发给对应学院负责人"""
import os

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import conf

BASE = conf.BASE
OUT = conf.OUT_DIR

SCHOOLS = {
    "机械学院": ("机械学院", "机械工程与智能制造学院"),
    "数据学院": ("人工智能学院", "计算机与人工智能学院（原数据学院）"),
}

thin = Side(style="thin", color="D9D9D9")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def style(cell, fill=None, bold=False, color="000000"):
    cell.font = Font(bold=bold, color=color)
    cell.border = BOX
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def title(ws, text, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, text)
    c.font = Font(size=14, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30


def dump(ws, header, rows, widths, start=3):
    for j, h in enumerate(header, 1):
        style(ws.cell(start - 1, j, h), "C00000", True, "FFFFFF")
    for i, r in enumerate(rows, start):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            c.border = BOX
            c.alignment = Alignment(vertical="center", wrap_text=True,
                                    horizontal="center" if j == 1 else "left")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ---------- 读现有的两张总表 ----------
src = openpyxl.load_workbook(os.path.join(OUT, "团员人数统计.xlsx"))
s_class = [list(r) for r in src["按班级"].iter_rows(min_row=2, values_only=True)]
s_issue = [list(r) for r in src["问题清单"].iter_rows(min_row=2, values_only=True)]

err = openpyxl.load_workbook(os.path.join(OUT, "团员花名册有误名单.xlsx"))
s_err = [list(r) for r in err["有误名单"].iter_rows(min_row=2, values_only=True)]

for fname, (key, full) in SCHOOLS.items():
    cls_rows = [r for r in s_class if str(r[0]).strip() == key]
    iss_rows = [r for r in s_issue if str(r[0]).strip() == key]
    err_rows = [r for r in s_err if str(r[0]).strip() == key]

    wb = openpyxl.Workbook()

    # ---- 人数统计 ----
    ws = wb.active
    ws.title = "人数统计"
    title(ws, f"{full}　2026级团员花名册核查", 12)
    ws.cell(2, 1, f"共 {len(cls_rows)} 个班　"
                  f"花名册 {sum(int(r[2] or 0) for r in cls_rows)} 人　"
                  f"系统内 {sum(int(r[3] or 0) for r in cls_rows)} 人　"
                  f"缺 {sum(int(r[4] or 0) for r in cls_rows)} 人").font = Font(size=10, color="C00000")
    dump(ws, ["专业班级", "花名册", "系统内", "系统里没有", "缺失人员名单",
              "系统里多出", "班级不符", "学号重复", "出现两次", "备注为空", "需处理"],
         [[r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10], r[11]] for r in cls_rows],
         [26, 9, 9, 11, 62, 11, 10, 10, 10, 10, 9])
    ws.freeze_panes = "A3"

    # ---- 问题清单 ----
    ws2 = wb.create_sheet("问题清单")
    ws2.cell(1, 1, f"{full}　需处理学生清单（共 {len(iss_rows)} 人）").font = Font(size=13, bold=True)
    dump(ws2, ["专业班级", "姓名", "学号", "问题说明"],
         [[r[1], r[2], r[3], r[4]] for r in iss_rows],
         [26, 12, 16, 58])
    ws2.freeze_panes = "A3"

    # ---- 有误名单（含籍贯/格式，备用） ----
    ws3 = wb.create_sheet("有误名单(备用)")
    ws3.cell(1, 1, f"{full}　逐人问题名单（含籍贯、时间格式，供后续使用）").font = Font(size=13, bold=True)
    dump(ws3, ["学号", "专业班级", "姓名", "籍贯", "性别", "出生年月",
               "入团时间", "入团地点", "备注", "问题类型"],
         [[r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10]] for r in err_rows],
         [14, 26, 10, 24, 6, 11, 11, 28, 10, 44])
    ws3.freeze_panes = "A3"

    path = os.path.join(OUT, f"{fname}-2026级团员核查.xlsx")
    wb.save(path)
    print(f"{fname}: {len(cls_rows)} 班 / 缺 {sum(int(r[4] or 0) for r in cls_rows)} 人 / "
          f"问题 {len(iss_rows)} 条 / 名单 {len(err_rows)} 行 -> {os.path.basename(path)}")

print("\n输出目录:", OUT)
