# -*- coding: utf-8 -*-
"""出两份文件

① 团员花名册有误名单.xlsx        提交版：问题类型只写名字，跟他们旧名单一致
② 团员花名册有误名单-复核版.xlsx  自用：带「应改为」、按错误形态分组、系统无分班

列顺序（两份相同）：
  学院 | 学号 | 专业班级 | 姓名 | 籍贯 | 性别 | 出生年月 | 入团时间 | 入团地点 | 备注 | 问题类型
"""
import os
import re
import collections

import importlib.util
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import conf

BASE = conf.BASE
OUT = conf.OUT_DIR
P_A = os.path.join(OUT, "核查结果-表格内问题.xlsx")
P_B = os.path.join(OUT, "对账结果.xlsx")
DST_MAIN = os.path.join(OUT, "团员花名册有误名单.xlsx")
DST_CHK = os.path.join(OUT, "团员花名册有误名单-复核版.xlsx")

HEAD = ["学院", "学号", "专业班级", "姓名", "籍贯", "性别", "出生年月",
        "入团时间", "入团地点", "备注", "问题类型"]
WID = [12, 14, 26, 10, 26, 6, 11, 11, 30, 11, 40]
thin = Side(style="thin", color="D9D9D9")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
# 团支部全称 -> 简称（沿用花名册自己的叫法）
_nm = importlib.util.spec_from_file_location("nm", os.path.join(BASE, "scripts", "namemap.py"))
_m = importlib.util.module_from_spec(_nm); _nm.loader.exec_module(_m)
org_short = _m.org_short


def read(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    return [list(r) + [None] * (12 - len(r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def cls_note(note):
    """班级不符：他们原来的写法是「智慧团建上为机制236，与表格不符」"""
    m = re.search(r"智慧团建在[“\"]([^”\"]+)[”\"]", note or "")
    return f"智慧团建上为{org_short(m.group(1))}，与表格不符" if m else "班级不符"


def compose_plain(t, note):
    """提交版：只保留问题名，括号说明全部去掉"""
    out = []
    for lab in [x for x in str(t or "").split("、") if x]:
        out.append(cls_note(note) if lab == "班级不符" else lab)
    return "、".join(dict.fromkeys(out))


def compose_full(t, note, me=""):
    """复核版：带上具体怎么错、该改成什么"""
    note = (note or "").strip()
    out = []
    for lab in [x for x in str(t or "").split("、") if x]:
        if lab == "籍贯有误":
            m = re.search(r"应(?:只到|改为|只写|改)[“\"]([^”\"]+)[”\"]", note)
            out.append(f"籍贯有误（应改为“{m.group(1)}”）" if m else "籍贯有误")
        elif lab == "班级不符":
            out.append(cls_note(note))
        elif lab == "学号重复":
            m = re.search(r"被\s*(.+?)\s*共用", note)
            others = [x for x in (m.group(1).replace("、", "/").split("/") if m else []) if x and x != me]
            out.append(f"学号重复（与{'/'.join(others)}共用）" if others else "学号重复")
        elif lab == "时间格式错误":
            m = re.search(r"(出生年月|入团时间)(为空|[^；]+)", note)
            if m:
                why = (m.group(2).replace("多写了日，只写到月", "多写了日")
                       .replace("格式不对（应为 2024.9 这种）", "格式不是 2024.9")
                       .replace("只写到年，缺月份", "只写到年"))
                out.append(f"时间格式错误（{m.group(1)}{why[:14]}）")
            else:
                out.append("时间格式错误")
        else:
            out.append(lab)
    return "、".join(dict.fromkeys(out))


def rowkey(r):
    cls = str(r[2] or "")
    m = re.match(r"(\D+)(\d+)(?:班)?", cls)
    return (0 if str(r[0]) == "机械学院" else 1,
            m.group(1) if m else cls,
            int(m.group(2)) if m and m.group(2) else 0,
            str(r[3] or ""))


def merge(rows):
    """同一人的多个问题并成一行"""
    merged, order = {}, []
    for r in rows:
        k = (str(r[3] or ""), str(r[1] or ""), str(r[2] or ""))
        if k not in merged:
            merged[k] = list(r)
            order.append(k)
        else:
            t0, t1 = str(merged[k][10] or ""), str(r[10] or "")
            merged[k][10] = "、".join(dict.fromkeys((t0 + "、" + t1).split("、")))
            for j in range(11):
                if merged[k][j] in (None, "") and r[j] not in (None, ""):
                    merged[k][j] = r[j]
    out = [merged[k] for k in order]
    out.sort(key=rowkey)
    return out


def write_list(wb, rows, title="有误名单"):
    ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = title
    for j, h in enumerate(HEAD, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="C00000")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[1].height = 24
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            if j == 2 and isinstance(v, (int, float)):
                v = str(int(v))
            c = ws.cell(i, j, v)
            c.border = BOX
            c.alignment = Alignment(vertical="center", wrap_text=(j == 11),
                                    horizontal="center" if j in (1, 6, 7, 8) else "left")
    for j, w in enumerate(WID, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{len(rows) + 1}"
    return ws


a = read(P_A)
b = read(P_B)

rows_plain = merge([[r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9],
                     compose_plain(r[10], r[11])] for r in a + b])
rows_full = merge([[r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9],
                    compose_full(r[10], r[11], r[3])] for r in a + b])


def add_summary(wb, rows):
    ws2 = wb.create_sheet("汇总")
    ws2["A1"] = "2026 级团员花名册核查汇总"
    ws2["A1"].font = Font(size=15, bold=True)
    ws2.merge_cells("A1:D1")
    ws2["A1"].alignment = Alignment(horizontal="center")
    ws2.row_dimensions[1].height = 30
    cnt = collections.Counter()
    for r in rows:
        for t in re.sub(r"（[^）]*）", "", str(r[10])).split("、"):
            if t:
                cnt[t] += 1
    ws2.append([])
    ws2.append(["问题类型", "人次"])
    for c in ws2[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center")
    for t, n in cnt.most_common():
        ws2.append([t, n])
    ws2.append(["合计", sum(cnt.values())])
    ws2.append([])
    ws2.append(["问题记录数", f"{len(rows)} 行（一人一行）"])
    ws2.column_dimensions["A"].width = 42
    ws2.column_dimensions["B"].width = 22

    ws3 = wb.create_sheet("按班级")
    ws3.append(["学院", "专业班级", "问题条数", "涉及学生"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    per = collections.defaultdict(list)
    for r in rows:
        per[(r[0], r[2])].append(r[3])
    for (col, cls), names in sorted(per.items(), key=lambda x: -len(x[1])):
        ws3.append([col, cls, len(names), "、".join(dict.fromkeys(n for n in names if n))])
    for j, w in enumerate([12, 26, 10, 70], 1):
        ws3.column_dimensions[get_column_letter(j)].width = w


# ---------- ① 提交版 ----------
wb = openpyxl.Workbook()
write_list(wb, rows_plain)
add_summary(wb, rows_plain)
wb.save(DST_MAIN)

# ---------- ② 复核版 ----------
wb2 = openpyxl.Workbook()
write_list(wb2, rows_full)
add_summary(wb2, rows_full)


def form_of(note):
    if "地级市名字可能写错" in note:
        return "④ 市名写错（疑似错别字）"
    if "多写了县级" in note:
        return "① 多写了县/区（删掉最后一段）"
    if "是县级" in note or "是区县" in note:
        return "② 县/区当成市（补上所属地级市）"
    if "市名重复" in note:
        return "③ 市名写了两遍"
    if "不是该省" in note:
        return "④ 市名写错/认不出"
    if "只写了省" in note:
        return "⑤ 只写了省，没写市"
    if "看不出省份" in note:
        return "⑥ 缺省名"
    return "⑦ 其它"


def suggest(note):
    m = re.search(r"应(?:只到|改为|只写|改)[“\"]([^”\"]+)[”\"]", note)
    return m.group(1) if m else ""


ws4 = wb2.create_sheet("籍贯待复核")
ws4.append(["形态", "学院", "专业班级", "姓名", "原籍贯", "建议改为"])
for c in ws4[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="BF8F00")
    c.alignment = Alignment(horizontal="center")
jg = []
for r in a:
    if "籍贯有误" not in str(r[10] or ""):
        continue
    note = str(r[11] or "")
    jg.append([form_of(note), r[0], r[2], r[3], r[4], suggest(note)])
jg.sort(key=lambda x: (x[0], str(x[2]), str(x[3])))
for row in jg:
    ws4.append(row)
for j, w in enumerate([26, 11, 24, 10, 26, 14], 1):
    ws4.column_dimensions[get_column_letter(j)].width = w
ws4.freeze_panes = "A2"
ws4.auto_filter.ref = f"A1:F{len(jg) + 1}"

ws5 = wb2.create_sheet("系统无待复核")
ws5.append(["学院", "专业班级", "姓名", "学号"])
for c in ws5[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="BF8F00")
    c.alignment = Alignment(horizontal="center")
miss = [[r[0], r[2], r[3], r[1]] for r in b if "智慧团建无，电子版有" in str(r[10] or "")]
miss.sort(key=lambda x: (str(x[0]), str(x[1]), str(x[2])))
for row in miss:
    ws5.append(row)
for j, w in enumerate([12, 26, 12, 14], 1):
    ws5.column_dimensions[get_column_letter(j)].width = w
ws5.freeze_panes = "A2"
ws5.auto_filter.ref = f"A1:D{len(miss) + 1}"

wb2.save(DST_CHK)

print("① 提交版:", os.path.basename(DST_MAIN), f"{len(rows_plain)} 行")
print("② 复核版:", os.path.basename(DST_CHK), f"{len(rows_full)} 行")
print("\n问题类型样例：")
for r in rows_plain[:3]:
    print("   ", r[2], r[3], "→", r[10])
print("\n籍贯形态：")
for f_, n in collections.Counter(x[0] for x in jg).most_common():
    print(f"   {n:4}  {f_}")
